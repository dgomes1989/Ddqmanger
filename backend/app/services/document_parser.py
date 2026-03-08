"""Parse Word, Excel, and PDF documents to extract raw text."""

import io

import docx
import openpyxl
import pdfplumber


def parse_docx(file_bytes: bytes) -> str:
    """Extract text from a Word document."""
    doc = docx.Document(io.BytesIO(file_bytes))
    paragraphs = []
    for para in doc.paragraphs:
        if para.text.strip():
            paragraphs.append(para.text.strip())

    # Also extract from tables
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))

    return "\n\n".join(paragraphs)


def parse_xlsx(file_bytes: bytes) -> str:
    """Extract text from an Excel workbook."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    sections = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            sections.append(f"[Sheet: {sheet_name}]\n" + "\n".join(rows))

    wb.close()
    return "\n\n".join(sections)


def parse_pdf(file_bytes: bytes) -> str:
    """Extract text from a PDF document."""
    pages = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text and text.strip():
                pages.append(text.strip())

            # Also try extracting tables
            tables = page.extract_tables()
            for table in tables:
                for row in table:
                    cells = [str(c).strip() for c in row if c and str(c).strip()]
                    if cells:
                        pages.append(" | ".join(cells))

    return "\n\n".join(pages)


def parse_document(file_bytes: bytes, file_type: str) -> str:
    """Parse a document based on its file type."""
    parsers = {
        "docx": parse_docx,
        "xlsx": parse_xlsx,
        "pdf": parse_pdf,
    }
    parser = parsers.get(file_type)
    if not parser:
        raise ValueError(f"Unsupported file type: {file_type}")
    return parser(file_bytes)
